"""File system tools for the resume assistant.

Every tool returns JSON-serialisable data and never raises: failures come back
as ``{"success": False, "error": "..."}`` so an LLM can read and react to them.

All paths are confined to a sandbox root (``FS_ROOT`` env var, default: the
current working directory) so a model cannot read or write outside it.
"""

from __future__ import annotations

import os
import re
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from pypdf import PdfReader

READABLE_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
WRITABLE_EXTENSIONS = {".txt", ".md", ".json", ".csv"}
MAX_CONTENT_CHARS = 20_000
MAX_MATCHES = 50
CONTEXT_CHARS = 80


class ToolError(Exception):
    """Expected, user-presentable failure inside a tool."""


def _root() -> Path:
    return Path(os.environ.get("FS_ROOT") or Path.cwd()).expanduser().resolve()


def _resolve(path: str) -> Path:
    """Resolve ``path`` (relative to the sandbox root) and ensure it stays inside."""
    root = _root()
    p = Path(path).expanduser()
    if not p.is_absolute():
        p = root / p
    p = p.resolve()
    if p != root and root not in p.parents:
        raise ToolError(f"Access denied: '{path}' is outside the allowed directory")
    return p


def _iso(ts: float) -> str:
    return datetime.fromtimestamp(ts, tz=timezone.utc).isoformat(timespec="seconds")


def _error(exc: Exception) -> dict:
    if isinstance(exc, ToolError):
        return {"success": False, "error": str(exc)}
    return {"success": False, "error": f"{type(exc).__name__}: {exc}"}


def _extract_text(path: Path) -> tuple[str, dict]:
    """Return (text, extra_metadata) for a supported file type."""
    if not path.exists():
        raise ToolError(f"File not found: {path.name}")
    if not path.is_file():
        raise ToolError(f"Not a file: {path.name}")
    ext = path.suffix.lower()
    if ext not in READABLE_EXTENSIONS:
        raise ToolError(
            f"Unsupported file type '{ext or '(none)'}'. "
            f"Supported: {', '.join(sorted(READABLE_EXTENSIONS))}"
        )

    if ext in {".txt", ".md"}:
        raw = path.read_bytes()
        try:
            return raw.decode("utf-8"), {}
        except UnicodeDecodeError:
            return raw.decode("latin-1"), {}

    if ext == ".pdf":
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ToolError("PDF is password-protected")
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages).strip(), {"page_count": len(pages)}

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip(), {}


def read_file(filepath: str) -> dict:
    """Read a resume (.pdf, .txt, .md, .docx) and return its text and metadata."""
    try:
        path = _resolve(filepath)
        text, extra = _extract_text(path)
        stat = path.stat()
        result = {
            "success": True,
            "filepath": str(path),
            "filename": path.name,
            "file_type": path.suffix.lower().lstrip("."),
            "content": text[:MAX_CONTENT_CHARS],
            "truncated": len(text) > MAX_CONTENT_CHARS,
            "metadata": {
                "size_bytes": stat.st_size,
                "modified": _iso(stat.st_mtime),
                "char_count": len(text),
                "word_count": len(text.split()),
                **extra,
            },
        }
        if not text:
            result["warning"] = "No text extracted (file may be empty or a scanned image)"
        return result
    except Exception as exc:  # noqa: BLE001 - tools must never raise
        return _error(exc)


def list_files(directory: str, extension: str | None = None) -> list:
    """List files in a directory, optionally filtered by extension (e.g. '.pdf').

    Returns a list of {name, path, size_bytes, modified}. On failure returns a
    one-item list ``[{"success": False, "error": ...}]`` to keep the list shape.
    """
    try:
        path = _resolve(directory)
        if not path.exists():
            raise ToolError(f"Directory not found: {directory}")
        if not path.is_dir():
            raise ToolError(f"Not a directory: {directory}")
        ext = None
        if extension:
            ext = extension.lower()
            if not ext.startswith("."):
                ext = "." + ext
        items = []
        for child in sorted(path.iterdir(), key=lambda c: c.name.lower()):
            if not child.is_file():
                continue
            if ext and child.suffix.lower() != ext:
                continue
            stat = child.stat()
            items.append(
                {
                    "name": child.name,
                    "path": str(child),
                    "size_bytes": stat.st_size,
                    "modified": _iso(stat.st_mtime),
                }
            )
        return items
    except Exception as exc:  # noqa: BLE001
        return [_error(exc)]


def write_file(filepath: str, content: str) -> dict:
    """Write text to a file, creating parent directories. Overwrites existing files."""
    try:
        path = _resolve(filepath)
        ext = path.suffix.lower()
        if ext not in WRITABLE_EXTENSIONS:
            raise ToolError(
                f"Cannot write '{ext or '(none)'}' files. "
                f"Allowed: {', '.join(sorted(WRITABLE_EXTENSIONS))}"
            )
        if path.is_dir():
            raise ToolError(f"Path is a directory: {filepath}")
        existed = path.exists()
        path.parent.mkdir(parents=True, exist_ok=True)
        data = content.encode("utf-8")
        # Atomic write: temp file in the same directory, then rename over target.
        fd, tmp = tempfile.mkstemp(dir=path.parent, prefix=".tmp-")
        try:
            with os.fdopen(fd, "wb") as fh:
                fh.write(data)
            os.replace(tmp, path)
        except BaseException:
            Path(tmp).unlink(missing_ok=True)
            raise
        return {
            "success": True,
            "filepath": str(path),
            "bytes_written": len(data),
            "overwrote": existed,
        }
    except Exception as exc:  # noqa: BLE001
        return _error(exc)


def search_in_file(filepath: str, keyword: str) -> dict:
    """Case-insensitive keyword search; returns each hit with surrounding context."""
    try:
        if not keyword or not keyword.strip():
            raise ToolError("Keyword must not be empty")
        path = _resolve(filepath)
        text, _ = _extract_text(path)
        matches = []
        total = 0
        for m in re.finditer(re.escape(keyword), text, re.IGNORECASE):
            total += 1
            if len(matches) >= MAX_MATCHES:
                continue
            start = max(0, m.start() - CONTEXT_CHARS)
            end = min(len(text), m.end() + CONTEXT_CHARS)
            context = " ".join(text[start:end].split())
            matches.append(
                {
                    "line_number": text.count("\n", 0, m.start()) + 1,
                    "match": m.group(0),
                    "context": ("..." if start > 0 else "")
                    + context
                    + ("..." if end < len(text) else ""),
                }
            )
        return {
            "success": True,
            "filepath": str(path),
            "keyword": keyword,
            "match_count": total,
            "matches": matches,
            "truncated": total > len(matches),
        }
    except Exception as exc:  # noqa: BLE001
        return _error(exc)
