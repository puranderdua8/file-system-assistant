import pytest
from docx import Document


def make_pdf(path, text):
    """Write a minimal single-page PDF whose text layer contains ``text``."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(out)


@pytest.fixture
def root(tmp_path, monkeypatch):
    monkeypatch.setenv("FS_ROOT", str(tmp_path))
    return tmp_path


@pytest.fixture
def resumes(root):
    d = root / "resumes"
    d.mkdir()
    (d / "alice.txt").write_text(
        "Alice Smith\nSenior engineer.\nSkills: Python, SQL, python testing.\n"
    )
    doc = Document()
    doc.add_paragraph("Bob Jones")
    doc.add_paragraph("Experienced in Java and Kotlin.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Languages"
    table.rows[0].cells[1].text = "Java, Go"
    doc.save(d / "bob.docx")
    make_pdf(d / "carol.PDF", "Carol Lee - Data scientist with Python experience")
    (d / "notes.csv").write_text("a,b\n")
    return d
