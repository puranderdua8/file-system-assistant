"""Text extraction for each supported file type.

To support a new type, write a ``parse_xxx(path) -> (text, metadata)`` function
and add it to ``PARSERS``. Parsers can assume the file exists; they raise
``ToolError`` only for type-specific failures.
"""

from __future__ import annotations

from pathlib import Path
from typing import Callable

from docx import Document
from pypdf import PdfReader

from errors import ToolError

Parsed = tuple[str, dict]


def parse_text(path: Path) -> Parsed:
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8"), {}
    except UnicodeDecodeError:
        return raw.decode("latin-1"), {}


def parse_pdf(path: Path) -> Parsed:
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ToolError("PDF is password-protected")
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip(), {"page_count": len(pages)}


def parse_docx(path: Path) -> Parsed:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip(), {}


PARSERS: dict[str, Callable[[Path], Parsed]] = {
    ".txt": parse_text,
    ".md": parse_text,
    ".pdf": parse_pdf,
    ".docx": parse_docx,
}
SUPPORTED_EXTENSIONS = frozenset(PARSERS)


def extract_text(path: Path) -> Parsed:
    """Validate ``path`` and return (text, extra_metadata) using the matching parser."""
    if not path.exists():
        raise ToolError(f"File not found: {path.name}")
    if not path.is_file():
        raise ToolError(f"Not a file: {path.name}")
    ext = path.suffix.lower()
    parser = PARSERS.get(ext)
    if parser is None:
        raise ToolError(
            f"Unsupported file type '{ext or '(none)'}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    return parser(path)
